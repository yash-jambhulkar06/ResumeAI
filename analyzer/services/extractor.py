"""
Resume text extraction service supporting PDF, DOCX, TXT, and images (OCR).
"""
import os
from pathlib import Path
from docx import Document
from PIL import Image
from pypdf import PdfReader
import pytesseract

MAX_UPLOAD_SIZE = 5 * 1024 * 1024  # 5 MB
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".jpg", ".jpeg", ".png"}
UPLOAD_ERROR = (
    "We couldn't read this resume. Please upload a valid PDF, DOCX, TXT, "
    "JPG or PNG file."
)


def extract_pdf_text(file) -> str:
    """Extract text from a PDF file with page-by-page extraction."""
    reader = PdfReader(file)
    text = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text.append(page_text)
    return "\n".join(text).strip()


def extract_docx_text(file) -> str:
    """Extract text from a DOCX file."""
    document = Document(file)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    # Also extract text from tables if present
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n".join(paragraphs).strip()


def extract_txt_text(file) -> str:
    """Extract text from a TXT file with tolerant UTF-8 decoding."""
    return file.read().decode("utf-8", errors="ignore").strip()


def extract_image_text(file) -> str:
    """Extract text from an image using Pillow and Tesseract OCR."""
    tesseract_cmd = os.getenv("TESSERACT_CMD")
    if not tesseract_cmd and os.name == "nt":
        default_path = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
        if default_path.exists():
            tesseract_cmd = str(default_path)
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    try:
        with Image.open(file) as image:
            return pytesseract.image_to_string(image).strip()
    except Exception as exc:
        raise ValueError(UPLOAD_ERROR) from exc


def extract_resume_text(file) -> str:
    """Validate an uploaded resume and extract its text in memory."""
    extension = Path(file.name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. Please upload PDF, DOCX, TXT, JPG or PNG."
        )
    if file.size > MAX_UPLOAD_SIZE:
        raise ValueError("File is too large. Please upload a file under 5 MB.")
    if file.size == 0:
        raise ValueError("This file is empty. Please upload a readable resume.")

    try:
        file.seek(0)
        if extension == ".pdf":
            text = extract_pdf_text(file)
        elif extension == ".docx":
            text = extract_docx_text(file)
        elif extension == ".txt":
            text = extract_txt_text(file)
        else:
            text = extract_image_text(file)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(UPLOAD_ERROR) from exc

    if not text:
        raise ValueError(
            "We couldn't find any readable text in this resume. "
            "If this is a scanned PDF, please upload a TXT, DOCX, or direct image instead."
        )
    return text
